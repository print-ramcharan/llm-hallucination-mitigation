"""DOCX document parser using python-docx."""

from __future__ import annotations

import io
from typing import Any

import docx

from src.ingestion.base import BaseParser
from src.ingestion.models import Document, DocumentElement, DocumentType


class DocxParser(BaseParser):
    """Extracts text, headings, tables, and metadata from DOCX files."""

    @property
    def supported_types(self) -> list[DocumentType]:
        return [DocumentType.DOCX]

    def can_parse(self, extension: str, mime_type: str | None = None) -> bool:
        ext = extension.lower().lstrip(".")
        if ext in ("docx", "doc"):
            return True
        if mime_type and "wordprocessingml" in mime_type.lower():
            return True
        return False

    def parse_bytes(
        self,
        content: bytes,
        filename: str = "document.docx",
        file_size_bytes: int = 0,
        **kwargs,
    ) -> Document:
        """Parse DOCX bytes into a standardized Document object."""
        if not content:
            raise ValueError("Cannot parse empty DOCX byte stream")

        doc_stream = io.BytesIO(content)
        doc = docx.Document(doc_stream)

        elements: list[DocumentElement] = []
        docx_metadata: dict[str, Any] = {}

        # Extract core properties if available
        try:
            core_props = doc.core_properties
            for attr in ("author", "title", "subject", "keywords", "created", "modified", "revision"):
                val = getattr(core_props, attr, None)
                if val:
                    docx_metadata[attr] = str(val)
        except Exception:
            pass

        current_heading: str | None = None
        el_idx = 0
        full_text_blocks: list[str] = []

        # Iterate over paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            is_heading = style_name.startswith("Heading") or "Title" in style_name
            element_type = "heading" if is_heading else "paragraph"

            if is_heading:
                current_heading = text

            elements.append(
                DocumentElement(
                    id=f"temp_el_{el_idx}",
                    element_type=element_type,
                    content=text,
                    section_title=current_heading,
                    metadata={"style": style_name},
                )
            )
            full_text_blocks.append(text)
            el_idx += 1

        # Iterate over tables and extract cell text cleanly
        for table_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                # Avoid empty rows
                if any(row_cells):
                    table_rows.append(" | ".join(row_cells))

            if table_rows:
                table_text = "\n".join(table_rows)
                elements.append(
                    DocumentElement(
                        id=f"temp_el_{el_idx}",
                        element_type="table",
                        content=table_text,
                        section_title=current_heading,
                        metadata={"table_index": table_idx, "rows": len(table_rows)},
                    )
                )
                full_text_blocks.append(table_text)
                el_idx += 1

        full_content = "\n\n".join(full_text_blocks)

        if kwargs.get("document_type"):
            docx_metadata["document_type"] = kwargs["document_type"]
        if "extra_metadata" in kwargs and isinstance(kwargs["extra_metadata"], dict):
            docx_metadata.update(kwargs["extra_metadata"])

        return Document.create(
            source_name=filename,
            file_type=DocumentType.DOCX,
            elements=elements,
            file_size_bytes=file_size_bytes or len(content),
            extra_metadata=docx_metadata,
            raw_content=full_content,
        )
