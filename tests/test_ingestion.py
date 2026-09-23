"""Unit and integration tests for the Ingestion Module."""

import io
from pathlib import Path

import docx
import fitz  # PyMuPDF
import pytest

from src.ingestion.docx import DocxParser
from src.ingestion.html import HtmlParser
from src.ingestion.markdown import MarkdownParser
from src.ingestion.models import (
    Document,
    DocumentElement,
    DocumentType,
)
from src.ingestion.pdf import PDFParser
from src.ingestion.registry import ParserRegistry, UnsupportedFormatError
from src.ingestion.store import DocumentStore
from src.ingestion.txt import TxtParser


class TestTxtParser:
    def test_parse_simple_txt(self):
        parser = TxtParser()
        content = b"Paragraph 1 line A.\nParagraph 1 line B.\n\nParagraph 2 text."
        doc = parser.parse_bytes(content, filename="sample.txt")

        assert doc.metadata.source_name == "sample.txt"
        assert doc.metadata.file_type == DocumentType.TXT
        assert len(doc.elements) == 2
        assert "Paragraph 1" in doc.elements[0].content
        assert "Paragraph 2" in doc.elements[1].content
        assert doc.metadata.word_count > 0
        assert doc.metadata.char_count > 0

    def test_parse_latin1_encoding(self):
        parser = TxtParser()
        content = "Café au lait\n\nCrème brûlée".encode("latin-1")
        doc = parser.parse_bytes(content, filename="french.txt")

        assert "Café" in doc.content or "Crème" in doc.content
        assert len(doc.elements) == 2

    def test_empty_txt_raises_error(self):
        parser = TxtParser()
        with pytest.raises(ValueError, match="Cannot parse empty"):
            parser.parse_bytes(b"", filename="empty.txt")


class TestMarkdownParser:
    def test_parse_markdown_with_hierarchy_and_code(self):
        parser = MarkdownParser()
        md_text = """---
title: Research Doc
author: Researcher
---

# Introduction to Context Degradation
Long contexts suffer from lost-in-the-middle degradation.

## Empirical Findings
Attention is concentrated at boundaries.

```python
def mitigate(query, docs):
    return rerank(docs)
```

### Conclusion
Multi-stage retrieval improves faithfulness.
"""
        doc = parser.parse_bytes(md_text.encode("utf-8"), filename="research.md")

        assert doc.metadata.file_type == DocumentType.MD
        assert doc.metadata.extra.get("title") == "Research Doc"
        assert doc.metadata.extra.get("author") == "Researcher"

        # Check elements
        headings = [el for el in doc.elements if el.element_type == "heading"]
        assert len(headings) >= 3
        assert any(el.metadata.get("title") == "Introduction to Context Degradation" for el in headings)

        code_blocks = [el for el in doc.elements if el.element_type == "code"]
        assert len(code_blocks) == 1
        assert "def mitigate" in code_blocks[0].content
        assert code_blocks[0].metadata.get("language") == "python"

    def test_empty_md_raises_error(self):
        parser = MarkdownParser()
        with pytest.raises(ValueError):
            parser.parse_bytes(b"", filename="empty.md")


class TestHtmlParser:
    def test_parse_html_strips_scripts_and_extracts_clean_text(self):
        parser = HtmlParser()
        html_content = """<!DOCTYPE html>
<html>
<head>
    <title>Hallucination Study</title>
    <meta name="description" content="Study on LLM hallucination" />
    <style>body { color: red; }</style>
    <script>console.log("malicious code");</script>
</head>
<body>
    <header><p>Site Header</p></header>
    <nav><a href="#">Link</a></nav>
    <h1>Evaluating Faithfulness</h1>
    <p>Faithfulness measures whether the answer is derived solely from the retrieved context.</p>
    <table>
        <tr><th>Metric</th><th>Score</th></tr>
        <tr><td>Precision</td><td>0.94</td></tr>
    </table>
    <footer><p>Copyright 2026</p></footer>
</body>
</html>"""
        doc = parser.parse_bytes(html_content.encode("utf-8"), filename="study.html")

        assert doc.metadata.file_type == DocumentType.HTML
        assert doc.metadata.extra.get("title") == "Hallucination Study"
        assert doc.metadata.extra.get("description") == "Study on LLM hallucination"

        # Scripts, nav, and footer must be removed
        assert "console.log" not in doc.content
        assert "Site Header" not in doc.content
        assert "Copyright 2026" not in doc.content

        # Content must be preserved
        assert "Evaluating Faithfulness" in doc.content
        assert "Faithfulness measures" in doc.content
        assert "Precision" in doc.content

        # Element validation
        h1_el = next((el for el in doc.elements if el.element_type == "heading"), None)
        assert h1_el is not None
        assert "Evaluating Faithfulness" in h1_el.content


class TestDocxParser:
    def test_parse_docx(self):
        parser = DocxParser()

        # Create an in-memory docx file
        doc_obj = docx.Document()
        doc_obj.add_heading("Context Optimization Overview", level=1)
        doc_obj.add_paragraph("This paragraph outlines RRF and cross-encoder reranking.")
        table = doc_obj.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Model"
        table.cell(0, 1).text = "Accuracy"
        table.cell(1, 0).text = "Hybrid RAG"
        table.cell(1, 1).text = "92.5%"

        bio = io.BytesIO()
        doc_obj.save(bio)
        docx_bytes = bio.getvalue()

        doc = parser.parse_bytes(docx_bytes, filename="report.docx")

        assert doc.metadata.file_type == DocumentType.DOCX
        assert "Context Optimization Overview" in doc.content
        assert "This paragraph outlines RRF" in doc.content
        assert "Hybrid RAG" in doc.content
        assert len(doc.elements) >= 3


class TestPDFParser:
    def test_parse_pdf(self):
        parser = PDFParser()

        # Create a simple in-memory 2-page PDF with PyMuPDF
        pdf_doc = fitz.open()
        page1 = pdf_doc.new_page()
        page1.insert_text((50, 72), "Page 1: Benchmark Dataset Analysis\nGrounding evidence prevents hallucination.")

        page2 = pdf_doc.new_page()
        page2.insert_text((50, 72), "Page 2: Lost-in-the-Middle Phenomenon\nRecall drops when evidence is placed centrally.")

        pdf_bytes = pdf_doc.write()
        pdf_doc.close()

        doc = parser.parse_bytes(pdf_bytes, filename="paper.pdf")

        assert doc.metadata.file_type == DocumentType.PDF
        assert doc.metadata.page_count == 2
        assert "Page 1: Benchmark Dataset Analysis" in doc.content
        assert "Lost-in-the-Middle Phenomenon" in doc.content

        # Elements should be tagged with page numbers
        p1_elements = [el for el in doc.elements if el.page_number == 1]
        p2_elements = [el for el in doc.elements if el.page_number == 2]
        assert len(p1_elements) >= 1
        assert len(p2_elements) >= 1


class TestParserRegistry:
    def test_registry_dispatches_correctly(self):
        registry = ParserRegistry()

        txt_doc = registry.parse_bytes(b"Hello text", filename="note.txt")
        assert txt_doc.metadata.file_type == DocumentType.TXT

        md_doc = registry.parse_bytes(b"# Note\nContent", filename="note.md")
        assert md_doc.metadata.file_type == DocumentType.MD

        with pytest.raises(UnsupportedFormatError):
            registry.parse_bytes(b"data", filename="file.unknown_xyz")

    def test_supported_extensions(self):
        registry = ParserRegistry()
        exts = registry.get_supported_extensions()
        for expected in [".pdf", ".docx", ".txt", ".md", ".html"]:
            assert expected in exts


class TestDocumentStore:
    def test_store_crud(self, tmp_path: Path):
        store = DocumentStore(storage_dir=tmp_path)

        # Create dummy doc
        doc = Document.create(
            source_name="test.txt",
            file_type=DocumentType.TXT,
            elements=[DocumentElement(id="el_0", content="Sample content")],
            raw_content="Sample content",
        )

        store.add(doc)

        retrieved = store.get(doc.id)
        assert retrieved is not None
        assert retrieved.id == doc.id
        assert retrieved.content == "Sample content"

        # List
        all_meta = store.list_all()
        assert len(all_meta) == 1
        assert all_meta[0].document_id == doc.id

        # Delete
        deleted = store.delete(doc.id)
        assert deleted is True
        assert store.get(doc.id) is None
        assert len(store.list_all()) == 0
